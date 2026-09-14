# -*- coding: utf-8 -*-
"""
由 manuscript/draft_content.py 生成投稿初稿。

输出:
  manuscript/Manuscript_draft.docx   双倍行距、连续行号，符合多数期刊的送审格式
  manuscript/Manuscript_draft.md     仓库内查阅与 diff
  manuscript/TODO_checklist.md       所有占位标记的清单，逐条勾掉才能投

刻意不生成参考文献表：正文中的 [REF-n: …] 需替换为真实文献，
生成一份带占位的参考文献表只会增加误投风险。
"""
import os
import re
import sys

sys.path.insert(0, "manuscript")
import draft_content as C  # noqa: E402

OUT = "manuscript"
TODO_RE = re.compile(r"【TODO:([^】]*)】")
REF_RE = re.compile(r"\[REF-(\d+):([^\]]*)\]")
PROV_RE = re.compile(r"《([^》]*)》")


def write_markdown():
    lines = [f"# {C.TITLE}", "",
             f"*Running head: {C.RUNNING_HEAD}*", "",
             "> **这是初稿，不可直接投稿。**",
             f"> 正文含 {len(TODO_RE.findall(raw_text()))} 处 `【TODO】` 占位、"
             f"{len(set(REF_RE.findall(raw_text())))} 处 `[REF-n]` 待补文献。",
             "> `《数字》` 标记的值来自预试验（5 折交叉验证、中位数填补、未做影像重标注），",
             "> 正式分析后必须更新。清单见 `TODO_checklist.md`。", "",
             "## Abstract", ""]
    for head, text in C.ABSTRACT:
        lines += [f"**{head}.** {text}", ""]
    lines += [f"**Keywords:** {C.KEYWORDS}", "", "---", ""]

    for head, paras in C.BODY:
        level = "##" if re.match(r"^\d+\.\s", head) else "###"
        lines += [f"{level} {head}", ""]
        for p in paras:
            lines += [p, ""]

    lines += ["---", "", "## Declarations", ""]
    for head, text in C.DECLARATIONS:
        lines += [f"**{head}.** {text}", ""]
    lines += ["## Tables", ""] + [f"- {t}" for t in C.TABLES] + ["", "## Figures", ""]
    lines += [f"- {f}" for f in C.FIGURES] + [""]
    lines += ["## References", "",
              "刻意留空。正文中的 `[REF-n: …]` 标注了需要哪一类文献，",
              "请补入真实引用后再建立本表——生成带占位的参考文献表只会增加误投风险。", ""]

    with open(f"{OUT}/Manuscript_draft.md", "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))


def raw_text():
    parts = [t for _, t in C.ABSTRACT]
    parts += [p for _, ps in C.BODY for p in ps]
    parts += [t for _, t in C.DECLARATIONS] + C.TABLES + C.FIGURES
    return "\n".join(parts)


def write_checklist():
    text = raw_text()
    todos = TODO_RE.findall(text)
    refs = sorted(set(REF_RE.findall(text)), key=lambda x: int(x[0]))
    provs = PROV_RE.findall(text)

    lines = ["# 投稿前待办清单", "",
             f"由 `scripts/build_manuscript.py` 自动生成。共 {len(todos)} 处占位、"
             f"{len(refs)} 处待补文献、{len(provs)} 处预试验数值。", "",
             "## 一、正式分析完成后必须替换的占位", ""]
    lines += [f"- [ ] {t.strip()}" for t in todos]

    lines += ["", "## 二、需要补入的真实文献", "",
              "**不要编造文献。** 每条标注了需要哪一类，请检索后填入真实引用。", ""]
    lines += [f"- [ ] `[REF-{n}]` {d.strip()}" for n, d in refs]

    lines += ["", "## 三、来自预试验、须以正式分析结果更新的数值", "",
              "这些值来自 5 折交叉验证 + 中位数填补，**未**做影像重标注、",
              "**未**做多重插补与 Bootstrap 校正，不能作为最终结果发表。", ""]
    seen = []
    for v in provs:
        if v not in seen:
            seen.append(v)
    lines += [f"- [ ] {v}" for v in seen]

    lines += ["", "## 四、投稿前的格式核对", "",
              "- [ ] 核实 J Pediatr Surg 现行 author guidelines："
              "字数上限、结构式摘要格式与词数、图表数量、参考文献格式、预印本政策",
              "- [ ] TRIPOD+AI 清单逐条填写，作为补充材料",
              "- [ ] 与磁性异物论文逐项核对数字口径（队列定义、磁体例数、穿孔数、伦理批件号）",
              "- [ ] 按方案 §七 排查 Fig 2 与 Table 4 是否与磁性异物论文重复",
              "- [ ] 投稿信中主动披露同队列的在审稿件",
              "- [ ] 确认补充材料不含任何原始自由文本", ""]

    with open(f"{OUT}/TODO_checklist.md", "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    return len(todos), len(refs), len(seen)


def write_docx():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    style.paragraph_format.space_after = Pt(0)

    # 送审稿惯例：连续行号，便于审稿人逐行指认
    sect = doc.sections[0]
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "360")
    sect._sectPr.append(ln)

    def para(text, bold=False, italic=False, size=12, align=None, after=0):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        run.font.size = Pt(size)
        if align:
            p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        return p

    para(C.TITLE, bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    para(f"Running head: {C.RUNNING_HEAD}", italic=True, size=10,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    para("【TODO: 作者姓名、单位、通讯作者联系方式】", size=10,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

    para("Abstract", bold=True, size=13, after=4)
    for head, text in C.ABSTRACT:
        p = doc.add_paragraph()
        r = p.add_run(f"{head}. ")
        r.bold = True
        p.add_run(text)
        p.paragraph_format.space_after = Pt(4)
    para(f"Keywords: {C.KEYWORDS}", italic=True, after=14)

    for head, paras in C.BODY:
        top = bool(re.match(r"^\d+\.\s", head))
        para(head, bold=True, size=13 if top else 12, after=4)
        for p in paras:
            para(p, after=4)

    doc.add_page_break()
    para("Declarations", bold=True, size=13, after=4)
    for head, text in C.DECLARATIONS:
        p = doc.add_paragraph()
        r = p.add_run(f"{head}. ")
        r.bold = True
        p.add_run(text)
        p.paragraph_format.space_after = Pt(4)

    para("", after=8)
    para("Tables", bold=True, size=13, after=4)
    for t in C.TABLES:
        para(t, after=2)
    para("", after=8)
    para("Figures", bold=True, size=13, after=4)
    for f in C.FIGURES:
        para(f, after=2)

    para("", after=8)
    para("References", bold=True, size=13, after=4)
    para("刻意留空——正文中的 [REF-n: …] 标注了需要哪一类文献，"
         "请补入真实引用后再建立本表。", italic=True, size=10)

    doc.save(f"{OUT}/Manuscript_draft.docx")


def main():
    os.makedirs(OUT, exist_ok=True)
    write_markdown()
    n_todo, n_ref, n_prov = write_checklist()
    write_docx()
    text = raw_text()
    words = len(re.sub(r"【[^】]*】|\[REF-[^\]]*\]|《|》", " ", text).split())
    print(f"已生成 {OUT}/Manuscript_draft.docx 与 .md")
    print(f"  正文约 {words} 词（不计占位标记）")
    print(f"  {n_todo} 处 TODO、{n_ref} 处待补文献、{n_prov} 处预试验数值")
    print(f"  清单见 {OUT}/TODO_checklist.md")


if __name__ == "__main__":
    main()

# -*- coding: utf-8 -*-
"""
从核验表单抽出双人子集，生成第二位核验者用的独立副本。

用法:
    python3 scripts/make_second_reader_form.py annotation/异物类型核验表单.xlsx

为什么单独做一份：Cohen's Kappa 只需要两人都读过的那 100 例，让第二位医师
把 411 例全读一遍是白费两小时。更要紧的是**独立性**——两人若在同一个文件上
先后填写，后填者会看到前者的判读，Kappa 就失去意义。故第二份必须是单独文件，
且不含第一位的任何填写内容。

保留原表的下拉校验、列宽、冻结窗格与盲法设置（不含治疗方式与结局），
只换掉说明页里关于分组例数的那几句。
"""
import os
import shutil
import sys

import openpyxl


def main():
    if len(sys.argv) < 2:
        sys.exit("用法: python3 scripts/make_second_reader_form.py <核验表单.xlsx>")
    src = sys.argv[1]
    if not os.path.exists(src):
        sys.exit(f"找不到 {src}")

    root, ext = os.path.splitext(src)
    dst = f"{root}_第二核验者{ext}"
    shutil.copy(src, dst)

    wb = openpyxl.load_workbook(dst)
    ws = wb["核验表"]
    hdr = [c.value for c in ws[1]]
    col_dual = hdr.index("双人核验") + 1

    # 从后往前删，否则行号会在删除过程中错位
    keep = 0
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, col_dual).value == "是":
            keep += 1
        else:
            ws.delete_rows(r)

    # 序号重排，便于第二位核验者与第一位对照进度（编号列仍是配对依据）
    col_no = hdr.index("序号") + 1
    for i, r in enumerate(range(2, ws.max_row + 1), start=1):
        ws.cell(r, col_no).value = i

    # delete_rows 不会收缩数据校验的作用域，重设到实际行数，
    # 否则下拉会一直延伸到原来的第 412 行
    from openpyxl.utils import get_column_letter

    for dv in ws.data_validations.dataValidation:
        letters = sorted({get_column_letter(c)
                          for rng in dv.sqref.ranges
                          for c in range(rng.min_col, rng.max_col + 1)})
        dv.sqref = " ".join(f"{L}2:{L}{ws.max_row}" for L in letters)

    note = wb["使用说明"]
    for row in note.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and "411" in cell.value:
                cell.value = (f"本份为**第二核验者专用**，只含双人核验子集 {keep} 例，"
                              "用于计算 Cohen's Kappa。请独立填写，"
                              "切勿参照第一位核验者的结果。预计 40 分钟。")

    wb.save(dst)
    print(f"已生成 {dst}")
    print(f"  保留 {keep} 行（双人核验子集）")
    print(f"  校验作用域已收缩至第 {ws.max_row} 行")


if __name__ == "__main__":
    main()

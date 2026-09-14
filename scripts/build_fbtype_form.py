# -*- coding: utf-8 -*-
"""
生成异物类型人工核验表单 annotation/异物类型核验表单.xlsx

核验分三组，各自回答一个不同的问题：
  A 未归类   正则未命中任一类别，需人工读现病史补全       —— 补全覆盖率
  B 多类命中  同时命中 ≥2 类，需人工判定主类               —— 消解歧义
  C 抽检     从单类命中中随机抽取，人工独立判读后与正则比对 —— 正则的准确率

与 X 线重标注不同，本任务的信息源就是文本本身，故表单直接呈现主诉与现病史，
无需另开病历系统。为保持盲法，表单不含治疗方式、结局与初步诊断。

投稿需要的两个指标：
  正则 vs 人工 的准确率(C 组)  —— 审稿人问「自动抽取准不准」
  人工 vs 人工 的 Kappa(双人子集) —— 审稿人问「判读一不一致」
"""
import os
import re
import sys

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, "scripts")
from build_cohort import FB_TYPES  # noqa: E402

XLSX = "10年消化道异物原始数据.xlsx"
OUTDIR = "annotation"
SEED = 20260914
N_SPOTCHECK = 200
N_DOUBLE_READ = 100

BASE = "Arial"
INK = "1F2937"
HDR_REF = "D9E2E5"
HDR_REQ = "FFE699"
BAND = "F5F8F8"
RULE = Side(style="thin", color="B8C4C6")

LABELS = {
    "fb_battery": "纽扣电池", "fb_magnet": "磁性异物", "fb_sharp": "尖锐异物",
    "fb_coin": "硬币或金属钝物", "fb_round": "圆球状钝物", "fb_pit": "果核",
    "fb_long": "长条状异物", "fb_plastic": "塑料或玩具", "fb_mercury": "水银或体温计",
}
CHOICES = list(LABELS.values()) + ["食物团块", "其他", "原文未说明", "非异物病例"]

FIELDS = [
    ("核验者", None, 10, "姓名缩写"),
    ("核验日期", None, 12, "YYYY-MM-DD"),
    ("异物类型", ",".join(CHOICES), 16, "据现病史判定的主类；一例只填一个"),
    ("次要类型", ",".join(CHOICES), 16, "确有第二类异物时填，否则留空"),
    ("原文词语", None, 18, "现病史中指明类型的原词，如「巴克球」"),
    ("数目", "1,2,3-5,≥6,不能确定", 10, "原文所述枚数"),
    ("原文所述大小", None, 16, "现病史若写了尺寸就照抄，如「直径约1cm」；未写则留空"),
    ("备注", None, 24, "存疑之处"),
]

REF_ORDER = ["序号", "科研就诊编号", "核验组", "正则判定", "性别", "年龄_岁", "主诉", "现病史节选"]
REF_WIDTHS = {"序号": 7, "科研就诊编号": 13, "核验组": 20, "正则判定": 18,
              "性别": 7, "年龄_岁": 8, "主诉": 20, "现病史节选": 62}


def excerpt(text, width=220):
    """截取现病史中与异物描述相关的一段，优先围绕「可能为/误食」展开。"""
    t = re.sub(r"\s+", " ", str(text)).strip()
    m = re.search(r"(?:可能为|考虑为|误[食服吞]|吞[食服]|不慎)", t)
    if m:
        start = max(0, m.start() - 40)
        return ("…" if start else "") + t[start:start + width] + ("…" if len(t) > start + width else "")
    return t[:width] + ("…" if len(t) > width else "")


def build():
    os.makedirs(OUTDIR, exist_ok=True)
    base = pd.read_excel(XLSX, sheet_name="病案首页基本信息")
    adm = pd.read_excel(XLSX, sheet_name="儿科入院记录")
    df = base.merge(adm, on=["科研患者编号", "科研就诊编号"], how="left")

    history = (df["主诉"].fillna("") + " " + df["现病史"].fillna("") + " "
               + df["门（急）诊诊断名称"].fillna("") + " " + df["初步诊断"].fillna("")).astype(str)

    cats = [k for k in FB_TYPES if k != "fb_multiple"]
    hits = pd.DataFrame({k: history.str.contains(FB_TYPES[k], regex=True) for k in cats})
    n_hit = hits.sum(axis=1)

    df["正则判定"] = [
        "、".join(LABELS[c] for c in cats if row[c]) or "（未归类）"
        for _, row in hits.iterrows()
    ]

    rng = np.random.default_rng(SEED)
    single = df.index[n_hit == 1].to_numpy()
    spot = set(rng.choice(single, size=min(N_SPOTCHECK, len(single)), replace=False))

    group = pd.Series("", index=df.index, dtype=object)
    group[n_hit == 0] = "A 未归类 · 需补全"
    group[n_hit >= 2] = "B 多类命中 · 需定主类"
    group[df.index.isin(spot)] = "C 抽检 · 核对正则"
    sel = group != ""

    work = df.loc[sel].copy()
    work["核验组"] = group[sel]
    work["年龄_岁"] = work["年龄（岁）"].round(1)
    work["现病史节选"] = work["现病史"].map(excerpt)
    work["主诉"] = work["主诉"].fillna("").astype(str).str.slice(0, 30)

    order = {"A 未归类 · 需补全": 0, "B 多类命中 · 需定主类": 1, "C 抽检 · 核对正则": 2}
    work = work.sort_values("核验组", key=lambda c: c.map(order)).reset_index(drop=True)
    work.insert(0, "序号", range(1, len(work) + 1))

    dbl = set(rng.choice(work.index.to_numpy(),
                         size=min(N_DOUBLE_READ, len(work)), replace=False))
    work["双人核验"] = np.where(work.index.isin(dbl), "是", "")

    cols = REF_ORDER + ["双人核验"]
    work[cols].to_csv(f"{OUTDIR}/fbtype_worklist.csv", index=False, encoding="utf-8-sig")

    wb = Workbook()
    ws = wb.active
    ws.title = "核验表"
    ws.append(cols + [f[0] for f in FIELDS])
    for row in work[cols].fillna("").itertuples(index=False):
        ws.append(list(row) + [None] * len(FIELDS))

    n_ref = len(cols)
    for i in range(1, ws.max_column + 1):
        c = ws.cell(row=1, column=i)
        c.fill = PatternFill("solid", fgColor=HDR_REF if i <= n_ref else HDR_REQ)
        c.font = Font(name=BASE, bold=True, size=10, color=INK)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=Side(style="medium", color="6B7F84"))
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = ws.cell(row=2, column=n_ref + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{len(work) + 1}"

    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = REF_WIDTHS.get(c, 11)
    for j, (_n, _o, width, _h) in enumerate(FIELDS):
        ws.column_dimensions[get_column_letter(n_ref + 1 + j)].width = width

    for j, (name, opts, _w, _h) in enumerate(FIELDS):
        if not opts:
            continue
        col = get_column_letter(n_ref + 1 + j)
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True,
                            showErrorMessage=True, errorTitle="取值不在允许范围",
                            error="请从下拉列表中选择。")
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{len(work) + 1}")

    band = PatternFill("solid", fgColor=BAND)
    wrap_cols = {cols.index("现病史节选") + 1, cols.index("主诉") + 1}
    for r in range(2, len(work) + 2):
        for i in range(1, ws.max_column + 1):
            c = ws.cell(row=r, column=i)
            c.font = Font(name=BASE, size=9, color=INK)
            c.alignment = Alignment(horizontal="left" if i in wrap_cols else "center",
                                    vertical="top", wrap_text=i in wrap_cols)
            c.border = Border(bottom=RULE)
            if r % 2 == 0 and i <= n_ref:
                c.fill = band
        ws.row_dimensions[r].height = 42

    # ---- 使用说明 ----
    info = wb.create_sheet("使用说明", 0)
    info.sheet_view.showGridLines = False
    info.column_dimensions["A"].width = 3
    info.column_dimensions["B"].width = 22
    info.column_dimensions["C"].width = 86

    def put(cell, text, size=10, bold=False, color=INK, wrap=False):
        c = info[cell]
        c.value = text
        c.font = Font(name=BASE, size=size, bold=bold, color=color)
        c.alignment = Alignment(vertical="top", wrap_text=wrap)

    put("B2", "异物类型人工核验", 15, True)
    put("B3", "三分类决策模型配套数据质控", 10, color="5A6B72")
    counts = work["核验组"].value_counts()
    rows = [
        ("为什么要做", "自动抽取的正则词典已按 1238 份现病史实测优化，未归类率由 22.7% 降至 14.9%，"
                       "但仍需人工核验：① 补全未归类者；② 判定多类命中者的主类；"
                       "③ 抽检以估计正则的准确率——投稿时审稿人必问自由文本抽取的可靠性。"),
        ("", ""),
        ("A 未归类 · 需补全", f"{counts.get('A 未归类 · 需补全', 0)} 例。"
                              "读现病史节选补填类型；原文确实没写就选「原文未说明」，不要猜。"),
        ("B 多类命中 · 需定主类", f"{counts.get('B 多类命中 · 需定主类', 0)} 例。"
                                  "正则同时命中多类，请判定主类；确有第二类异物时填「次要类型」。"),
        ("C 抽检 · 核对正则", f"{counts.get('C 抽检 · 核对正则', 0)} 例。"
                              "**先不看「正则判定」列**，自己读完再填，否则准确率会被高估。"),
        ("", ""),
        ("双人核验", f"{int((work['双人核验'] == '是').sum())} 例标为「是」，"
                     "由两人各填一份独立文件，用于 Cohen's Kappa。切勿互相参照。"),
        ("总行数", f"{len(work)} 例，预计 2–3 小时。"),
        ("", ""),
        ("怎么填", "只填黄色表头的列，灰蓝列是参考信息勿改。「异物类型」一例只填一个主类。"
                   "「原文词语」抄下现病史中指明类型的原词（如「巴克球」「枣胡」），"
                   "这些词会回补进词典，让下次抽取更准。"),
        ("关于尺寸", "只有 1.8%(22/1238) 的现病史写了具体尺寸，故本表不设测量字段——"
                      "尺寸以 X 线实测为准，见 annotation/X线重标注表单.xlsx 的长径/短径列。"
                      "「原文所述大小」仅为兜住这 1.8%，未写就留空，不要据描述估算。"),
        ("盲法", "表单刻意不含治疗方式、结局与初步诊断，避免因知道结局而改变判读。"),
        ("完成之后", "运行 python3 scripts/ingest_fbtype.py 表单.xlsx [第二位核验者.xlsx]"),
    ]
    r = 5
    for a, bb in rows:
        if a:
            put(f"B{r}", a, 10, True)
            put(f"C{r}", bb, 10, wrap=True)
            info.row_dimensions[r].height = 30 if len(bb) > 90 else 16
        r += 1

    out = f"{OUTDIR}/异物类型核验表单.xlsx"
    wb.save(out)
    print(f"已生成 {out}")
    print(work["核验组"].value_counts().to_string())
    print(f"双人核验子集: {int((work['双人核验'] == '是').sum())}")
    print(f"总计 {len(work)} 例（全队列 {len(df)} 例的 {len(work) / len(df) * 100:.1f}%）")


if __name__ == "__main__":
    build()

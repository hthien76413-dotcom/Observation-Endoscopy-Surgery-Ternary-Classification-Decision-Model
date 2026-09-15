# -*- coding: utf-8 -*-
"""
生成 X 线重标注表单 annotation/X线重标注表单.xlsx

来源: annotation/worklist.csv 与 annotation/original_reports.csv
      （由 scripts/build_worklist.py 生成）

工作表:
  使用说明   填写规则、字段速查、示例行
  标注表     主表，1073 行，参考列 + 标注列（下拉校验）
  代码本     每个字段的取值与定义
  原始报告   原放射科报告全文，按序号查阅（刻意与标注表分开，避免先入为主）
  进度       完成度与质控统计（公式，随填写自动更新）
"""
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = "Arial"
INK = "1F2937"
HDR_REF = "D9E2E5"      # 参考列表头：灰蓝
HDR_REQ = "FFE699"      # 必填列表头：黄
HDR_OPT = "E8EDD8"      # 选填列表头：浅绿
BAND = "F5F8F8"
RULE = Side(style="thin", color="B8C4C6")

# 标注字段: (列名, 必填?, 下拉选项 or None, 列宽, 说明)
FIELDS = [
    ("阅片者", True, None, 10, "填写者姓名缩写，如 ZS"),
    ("阅片日期", True, None, 12, "YYYY-MM-DD"),
    ("异物可见", True, "是,否,不确定", 10, "首次片上是否见不透 X 线异物"),
    ("部位", True,
     "食管上段,食管中段,食管下段,胃,十二指肠,小肠,结肠,直肠,多处,不能确定,无异物", 13,
     "异物所在消化道节段；多枚分布于不同节段填「多处」"),
    ("数目", True, "1,2,3-5,≥6,不能确定", 10, "可见异物枚数分档"),
    ("数目_精确", False, None, 10, "能准确计数时填整数，否则留空"),
    ("长径_mm", True, None, 10, "最大一枚异物的最长径，毫米，整数"),
    ("短径_mm", True, None, 10, "同一枚异物垂直于长径的宽度，毫米，整数"),
    ("成串成团", True, "是,否,不适用", 11, "多枚异物是否相互吸附成串或成团；单枚填「不适用」"),
    ("合并其他金属", True, "是,否,不适用", 12, "单枚磁体是否与另一金属异物并存；非磁性异物填「不适用」"),
    ("形态", False, "圆盘状,球形,条状,不规则,其他,不能确定", 11, "最大一枚的形态"),
    ("双环征", False, "是,否,不确定,不适用", 10, "圆盘状异物的双环／晕环征，提示纽扣电池"),
    ("影像判断类型", False,
     "硬币样,纽扣电池样,磁性球形,尖锐金属,其他金属,非金属但可见,不能判断", 14,
     "仅据影像判断，不参考病史"),
    ("膈下游离气体", False, "是,否", 13, "立位片膈下新月形透亮影"),
    ("肠梗阻征象", False, "是,否", 12, "肠管扩张伴气液平"),
    ("备注", False, None, 26, "异常情况、与报告不符之处等"),
]

REF_ORDER = ["序号", "科研就诊编号", "性别", "年龄_岁", "入院时间", "检查时间",
             "距入院_h", "X线次数", "报告名称", "主诉", "片子时相", "双人阅片", "需特别处理"]

REF_WIDTHS = {
    "序号": 7, "科研就诊编号": 13, "性别": 7, "年龄_岁": 8, "入院时间": 16,
    "检查时间": 16, "距入院_h": 10, "X线次数": 9, "报告名称": 20, "主诉": 18,
    "片子时相": 14, "双人阅片": 10, "需特别处理": 34,
}


def style_header(ws, row, ncols_ref):
    for i in range(1, ws.max_column + 1):
        c = ws.cell(row=row, column=i)
        if i <= ncols_ref:
            fill = HDR_REF
        else:
            fill = HDR_REQ if FIELDS[i - ncols_ref - 1][1] else HDR_OPT
        c.fill = PatternFill("solid", fgColor=fill)
        c.font = Font(name=BASE, bold=True, size=10, color=INK)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=Side(style="medium", color="6B7F84"))


def build_instructions(ws, n_rows, counts):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 17
    for col in "CDEFGH":
        ws.column_dimensions[col].width = 15
    ws.column_dimensions["I"].width = 40

    def put(cell, text, size=10, bold=False, color=INK, wrap=False):
        c = ws[cell]
        c.value = text
        c.font = Font(name=BASE, size=size, bold=bold, color=color)
        c.alignment = Alignment(vertical="center", wrap_text=wrap)
        return c

    put("B2", "儿童消化道异物 · X 线重标注表单", 16, True)
    put("B3", "三分类决策模型（观察 / 内镜 / 手术）配套数据采集", 10, color="5A6B72")

    put("B5", "为什么要做这件事", 12, True)
    put("B6", "原放射科报告的定位文本无法使用：1073 例中 557 例仅写「消化道不透 X 光异物」，"
              "未定位到具体节段。结果是影像变量对模型几乎无贡献（宏平均 AUC 仅 +0.002）。"
              "而异物在食管／胃 还是 幽门以远，本该是「内镜 vs 观察」的第一生理依据。",
        10, wrap=True)
    ws.merge_cells("B6:I9")

    put("B11", "工作量与分配", 12, True)
    rows = [
        ("总行数", f"{n_rows} 行", "每行 = 一次住院的首次 X 线"),
        ("术前片", f"{counts.get('术前片', 0)} 行", "拍摄早于手术开始，可直接用于建模"),
        ("无手术", f"{counts.get('无手术', 0)} 行", "本次住院无手术，全部可用"),
        ("同日-时序不明", f"{counts.get('同日-时序不明', 0)} 行", "需在 PACS 中核对拍摄时刻是否早于手术"),
        ("决策后片", f"{counts.get('决策后片', 0)} 行", "首次片已晚于手术开始，本次住院无术前片；仍请标注，但建模时剔除"),
        ("双人阅片", "200 行", "标注为「是」的行由两位阅片者独立完成，用于计算 Kappa"),
        ("预计耗时", "约 9 小时", "每例约 30 秒"),
    ]
    r = 12
    for a, bb, cc in rows:
        put(f"B{r}", a, 10, True)
        put(f"C{r}", bb, 10)
        put(f"E{r}", cc, 9, color="5A6B72")
        r += 1

    put("B21", "怎么填", 12, True)
    tips = [
        "1.  只填【黄色表头】和【浅绿表头】的列。灰蓝表头是参考信息，请勿修改。",
        "2.  黄色 = 必填，浅绿 = 选填但有分析价值，能填尽量填。",
        "3.  绝大多数列是下拉选择，点单元格右侧箭头即可，不要手输。",
        "4.  先看 PACS 原片形成自己的判断，再决定是否翻阅「原始报告」工作表。"
        "     原报告刻意与标注表分开，就是为了避免被原描述带偏。",
        "5.  「异物可见 = 否」时，部位填「无异物」，数目填「不能确定」，长径短径填 0，成串成团与合并其他金属填「不适用」。",
        "6.  长径与短径分开量：短径决定能否过幽门（限制因素是最小横截面），长径决定能否过十二指肠。"
    "     圆盘状异物两者相等。阈值在分析时套用，日后调整阈值不必重新阅片。",
    "7.  双人阅片行由两人各填一份独立文件，切勿互相参照，最后再合并算 Kappa。",
        "8.  拿不准就填「不确定」／「不能确定」，并在备注里写一句。不要猜。",
    ]
    r = 22
    for t in tips:
        put(f"B{r}", t, 10, wrap=True)
        ws.merge_cells(f"B{r}:I{r}")
        ws.row_dimensions[r].height = 26
        r += 1

    put(f"B{r + 1}", "示例行（格式参考，标注表中无此行）", 12, True)
    ex_head = ["阅片者", "阅片日期", "异物可见", "部位", "数目", "数目_精确",
               "长径_mm", "短径_mm", "成串成团", "合并其他金属", "形态", "影像判断类型", "备注"]
    ex_val = ["ZS", "2026-09-20", "是", "胃", "3-5", "4", "5", "5", "是", "否",
              "球形", "磁性球形", "四枚成串吸附，位于胃体"]
    hr = r + 2
    for i, (h, v) in enumerate(zip(ex_head, ex_val)):
        hc = ws.cell(row=hr, column=2 + i, value=h)
        hc.font = Font(name=BASE, bold=True, size=9, color=INK)
        hc.fill = PatternFill("solid", fgColor=HDR_REQ)
        hc.alignment = Alignment(horizontal="center", wrap_text=True)
        hc.border = Border(bottom=RULE)
        vc = ws.cell(row=hr + 1, column=2 + i, value=v)
        vc.font = Font(name=BASE, size=9)
        vc.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(2 + i)].width = max(11, len(h) + 4)
    ws.column_dimensions[get_column_letter(2 + len(ex_head) - 1)].width = 26

    put(f"B{hr + 3}", "完成后把文件交回，运行 scripts/ingest_annotation.py 即可并入分析数据集。",
        10, color="5A6B72")


def build_codebook(ws):
    ws.sheet_view.showGridLines = False
    heads = ["字段", "必填", "取值", "定义与判读要点"]
    widths = [15, 8, 46, 62]
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=BASE, bold=True, size=10, color=INK)
        c.fill = PatternFill("solid", fgColor=HDR_REF)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(bottom=Side(style="medium", color="6B7F84"))
        ws.column_dimensions[get_column_letter(i)].width = w

    defs = {
        "部位": "以横膈、幽门、回盲部为界。食管段按胸廓入口、气管隆凸再分上／中／下段；"
                "不能分段时填「不能确定」。胃与幽门以远的区分最关键——它直接决定内镜可及性。",
        "数目": "按可见枚数分档。成串吸附的磁珠按实际枚数计，不按串计。",
        "长径_mm": "最大一枚异物的最长径。请用 PACS 测量工具实测，不要目测估计。"
                   "长径决定能否通过十二指肠 C 形襻。",
        "短径_mm": "同一枚异物垂直于长径的宽度。短径决定能否通过幽门——"
                   "限制因素是最小横截面，不是最长边。硬币等圆盘状异物长径=短径。"
                   "两径分开记录，阈值在分析时套用，日后调整阈值无需重新阅片。",
        "合并其他金属": "单枚磁体与另一枚金属异物并存时，两者可隔肠壁相吸，"
                        "风险等同于多枚磁体。这是容易被漏掉的一种情况，请专门判断。"
                        "非磁性异物填「不适用」。",
        "成串成团": "两枚及以上异物相互吸附、聚集成串或成团。磁性异物的关键征象，"
                    "与肠壁夹持穿孔直接相关。单枚异物填「不适用」。",
        "双环征": "圆盘状异物边缘的双环／晕环（halo）影，正位可见；侧位可见台阶征。"
                  "是纽扣电池与硬币的鉴别要点。非圆盘状异物填「不适用」。",
        "影像判断类型": "仅凭影像判断，刻意不参考病史。与病史所述类型的不一致本身是有价值的数据，"
                        "请如实填写，不要向病史靠拢。",
        "膈下游离气体": "立位胸腹片膈下新月形透亮影。卧位片无法评估时填「否」并在备注注明。",
        "肠梗阻征象": "肠管扩张伴气液平面。",
        "异物可见": "首次片上是否见明确的不透 X 线异物。阴性同样是信息——"
                    "透光性异物（塑料、木质、部分鱼骨）本就拍不到。",
    }
    r = 2
    for name, req, opts, _w, note in FIELDS:
        ws.cell(row=r, column=1, value=name).font = Font(name=BASE, bold=True, size=10)
        ws.cell(row=r, column=2, value="必填" if req else "选填").font = Font(
            name=BASE, size=10, color="9A5B0B" if req else "5A6B72")
        ws.cell(row=r, column=3, value=opts.replace(",", " / ") if opts else "自由填写")
        ws.cell(row=r, column=4, value=defs.get(name, note))
        for i in range(1, 5):
            c = ws.cell(row=r, column=i)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=RULE)
            if i >= 3:
                c.font = Font(name=BASE, size=9)
        ws.row_dimensions[r].height = 42
        r += 1
    ws.freeze_panes = "A2"


def field_column(name, n_ref):
    """按字段名反查其在标注表中的列号，避免硬编码偏移。"""
    idx = [f[0] for f in FIELDS].index(name)
    return n_ref + 1 + idx


def build_progress(ws, n_rows, n_ref):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 44

    ws["B2"] = "标注进度与质控"
    ws["B2"].font = Font(name=BASE, size=14, bold=True, color=INK)
    ws["B3"] = "公式自动更新，无需手动维护"
    ws["B3"].font = Font(name=BASE, size=9, color="5A6B72")

    pos = get_column_letter(field_column("部位", n_ref))
    phase = get_column_letter(REF_ORDER.index("片子时相") + 1)
    dbl = get_column_letter(REF_ORDER.index("双人阅片") + 1)
    last = n_rows + 1

    rows = [
        ("总行数", f"=COUNTA(标注表!A2:A{last})", "工作清单总条目"),
        ("已标注", f"=COUNTA(标注表!{pos}2:{pos}{last})", "以「部位」列是否填写为准"),
        ("完成率", f"=IFERROR(C6/C5,0)", "已标注 ÷ 总行数"),
        ("", "", ""),
        ("术前片 已标注",
         f'=COUNTIFS(标注表!{phase}2:{phase}{last},"术前片",标注表!{pos}2:{pos}{last},"<>")',
         "可直接用于建模的部分"),
        ("无手术 已标注",
         f'=COUNTIFS(标注表!{phase}2:{phase}{last},"无手术",标注表!{pos}2:{pos}{last},"<>")',
         "同样可直接用于建模"),
        ("同日-时序不明 已标注",
         f'=COUNTIFS(标注表!{phase}2:{phase}{last},"同日-时序不明",标注表!{pos}2:{pos}{last},"<>")',
         "另需在 PACS 核对拍摄时刻"),
        ("", "", ""),
        ("双人阅片 应完成", f'=COUNTIF(标注表!{dbl}2:{dbl}{last},"是")', "Kappa 一致性检验子集"),
        ("双人阅片 已完成",
         f'=COUNTIFS(标注表!{dbl}2:{dbl}{last},"是",标注表!{pos}2:{pos}{last},"<>")',
         "本文件内的完成数；另一位阅片者请用独立文件"),
        ("", "", ""),
        ("定位到具体节段",
         f'=COUNTA(标注表!{pos}2:{pos}{last})'
         f'-COUNTIF(标注表!{pos}2:{pos}{last},"不能确定")'
         f'-COUNTIF(标注表!{pos}2:{pos}{last},"无异物")',
         "本次重标注的核心产出。原报告仅 516/1073 能定位，此数应显著高于它"),
    ]
    r = 5
    for label, formula, note in rows:
        if label:
            ws.cell(row=r, column=2, value=label).font = Font(name=BASE, size=10, bold=True)
            c = ws.cell(row=r, column=3, value=formula)
            c.font = Font(name=BASE, size=11)
            c.alignment = Alignment(horizontal="right")
            c.number_format = "0.0%" if label == "完成率" else "#,##0"
            ws.cell(row=r, column=4, value=note).font = Font(name=BASE, size=9, color="5A6B72")
            ws.cell(row=r, column=4).alignment = Alignment(vertical="center", wrap_text=True)
            for i in range(2, 5):
                ws.cell(row=r, column=i).border = Border(bottom=RULE)
        r += 1


def main():
    work = pd.read_csv("annotation/worklist.csv")
    reports = pd.read_csv("annotation/original_reports.csv")
    work = work.fillna("")
    assert list(work.columns) == REF_ORDER, f"工作清单列顺序与 REF_ORDER 不符: {list(work.columns)}"
    n = len(work)
    counts = work["片子时相"].value_counts().to_dict()

    wb = Workbook()
    build_instructions(wb.active, n, counts)
    wb.active.title = "使用说明"

    ws = wb.create_sheet("标注表")
    ref_cols = list(work.columns)
    ws.append(ref_cols + [f[0] for f in FIELDS])
    for row in work.itertuples(index=False):
        ws.append(list(row) + [None] * len(FIELDS))

    style_header(ws, 1, len(ref_cols))
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = ws.cell(row=2, column=len(ref_cols) + 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{n + 1}"

    for i, c in enumerate(ref_cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = REF_WIDTHS.get(c, 12)
    for j, (name, _req, _opts, width, _note) in enumerate(FIELDS):
        ws.column_dimensions[get_column_letter(len(ref_cols) + 1 + j)].width = width

    for j, (_name, _req, opts, _w, _note) in enumerate(FIELDS):
        if not opts:
            continue
        col = get_column_letter(len(ref_cols) + 1 + j)
        dv = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True,
                            showErrorMessage=True, errorTitle="取值不在允许范围",
                            error="请从下拉列表中选择。定义见「代码本」工作表。")
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n + 1}")

    band = PatternFill("solid", fgColor=BAND)
    for r in range(2, n + 2):
        for i in range(1, ws.max_column + 1):
            c = ws.cell(row=r, column=i)
            c.font = Font(name=BASE, size=9, color=INK)
            c.alignment = Alignment(horizontal="center" if i != len(ref_cols) else "left",
                                    vertical="center", wrap_text=(i == len(ref_cols)))
            c.border = Border(bottom=RULE)
            if r % 2 == 0 and i <= len(ref_cols):
                c.fill = band
        ws.row_dimensions[r].height = 20

    build_codebook(wb.create_sheet("代码本"))

    wr = wb.create_sheet("原始报告")
    wr.append(["序号", "科研就诊编号", "报告名称", "检查所见", "检查结论"])
    for row in reports.fillna("").itertuples(index=False):
        wr.append(list(row))
    for i, w in enumerate([7, 13, 22, 70, 46], start=1):
        wr.column_dimensions[get_column_letter(i)].width = w
    for i in range(1, 6):
        c = wr.cell(row=1, column=i)
        c.font = Font(name=BASE, bold=True, size=10, color=INK)
        c.fill = PatternFill("solid", fgColor=HDR_REF)
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=Side(style="medium", color="6B7F84"))
    for r in range(2, len(reports) + 2):
        for i in range(1, 6):
            c = wr.cell(row=r, column=i)
            c.font = Font(name=BASE, size=9)
            c.alignment = Alignment(vertical="top", wrap_text=(i >= 4))
            c.border = Border(bottom=RULE)
    wr.freeze_panes = "A2"
    wr.auto_filter.ref = f"A1:E{len(reports) + 1}"

    build_progress(wb.create_sheet("进度"), n, len(ref_cols))

    os.makedirs("annotation", exist_ok=True)
    out = "annotation/X线重标注表单.xlsx"
    wb.save(out)
    verify(out, n)
    print(f"已生成 {out}：{n} 行，{len(FIELDS)} 个标注字段")


def verify(path, n_rows):
    """LibreOffice 在部分环境不可用，无法用 recalc 验证；改为直接核对公式引用指向的表头。"""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws, prog = wb["标注表"], wb["进度"]
    n_ref = len(REF_ORDER)

    expect = {get_column_letter(field_column(f[0], n_ref)): f[0] for f in FIELDS}
    expect.update({get_column_letter(i + 1): c for i, c in enumerate(REF_ORDER)})
    for col, name in expect.items():
        got = ws[f"{col}1"].value
        assert got == name, f"列 {col} 表头应为 {name}，实为 {got}"

    import re as _re
    refs = 0
    for row in prog.iter_rows(min_col=3, max_col=3):
        v = row[0].value
        if not isinstance(v, str) or not v.startswith("="):
            continue
        for col, first, lastr in _re.findall(r"标注表!([A-Z]+)(\d+):[A-Z]+(\d+)", v):
            assert int(first) == 2 and int(lastr) == n_rows + 1, f"公式行范围有误: {v}"
            assert col in expect, f"公式引用了未知列 {col}: {v}"
            refs += 1
    assert refs >= 10, f"进度表公式引用数异常: {refs}"
    print(f"  自检通过：表头对位正确，进度表 {refs} 处区域引用均指向正确列与行范围")


if __name__ == "__main__":
    main()
